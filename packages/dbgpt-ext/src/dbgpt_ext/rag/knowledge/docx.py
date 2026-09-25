"""Docx Knowledge."""

from typing import Any, Dict, Iterator, List, Optional, Union

import docx
from docx.opc.oxml import parse_xml
from docx.opc.pkgreader import _SerializedRelationship, _SerializedRelationships

from dbgpt.core import Document
from dbgpt.rag.knowledge.base import (
    ChunkStrategy,
    DocumentType,
    Knowledge,
    KnowledgeType,
)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
# Content controls and custom XML wrap paragraphs, tables, rows and cells.
_WRAPPERS = frozenset({f"{_W}sdt", f"{_W}customXml"})
# Runs under these are not part of the text as the document reads: tracked
# deletions, text moved away, the ruby guide printed over its base text, text
# boxes (read as paragraphs of their own) and the fallback copy of a text box.
_SKIPPED_RUN_ANCESTORS = frozenset(
    {f"{_W}del", f"{_W}moveFrom", f"{_W}rt", f"{_W}txbxContent", _MC_FALLBACK}
)
# The run children python-docx's Run.text reads, each as the text it stands for.
_RUN_TEXT_TAGS = tuple(
    f"{_W}{tag}" for tag in ("br", "cr", "noBreakHyphen", "ptab", "t", "tab")
)


def _docx_blocks(container: Any) -> Iterator[str]:
    """Yield the text of each paragraph, table and text box in document order.

    ``Document.paragraphs`` lists only the paragraphs directly under the body,
    and ``Paragraph.text`` reads only the runs directly under a paragraph, so
    tables, text boxes, content controls, simple fields and tracked insertions
    never reached the loaded text.
    """
    for child in container:
        if child.tag == f"{_W}p":
            yield _paragraph_text(child)
            for text_box in _text_boxes(child):
                yield from _docx_blocks(text_box)
        elif child.tag == f"{_W}tbl":
            yield _table_text(child)
        elif child.tag in _WRAPPERS:
            yield from _docx_blocks(_wrapped_content(child))


def _paragraph_text(paragraph: Any) -> str:
    # Read the text elements themselves in document order, so a ruby base nested
    # inside a run stays between the text before and after it.
    return "".join(
        str(element)
        for element in paragraph.iter(*_RUN_TEXT_TAGS)
        if element.getparent().tag == f"{_W}r"
        and not _has_ancestor(element, paragraph, _SKIPPED_RUN_ANCESTORS)
    )


def _text_boxes(paragraph: Any) -> Iterator[Any]:
    for text_box in paragraph.iter(f"{_W}txbxContent"):
        # A nested text box is read with the box around it, and a fallback copy
        # repeats the real one.
        if not _has_ancestor(text_box, paragraph, _SKIPPED_RUN_ANCESTORS):
            yield text_box


def _table_text(table: Any) -> str:
    rows = []
    for row in _children(table, "tr"):
        # A tracked row or cell deletion is marked in its properties rather than
        # around its content.
        if _is_marked(row, "trPr", "del"):
            continue
        cells = (
            " ".join(block for block in _docx_blocks(cell) if block).replace("\n", " ")
            for cell in _children(row, "tc")
            if not _is_marked(cell, "tcPr", "cellDel")
        )
        rows.append(" | ".join(cells))
    return "\n".join(rows)


def _is_marked(element: Any, properties: str, marker: str) -> bool:
    found = element.find(f"{_W}{properties}")
    return found is not None and found.find(f"{_W}{marker}") is not None


def _children(parent: Any, tag: str) -> Iterator[Any]:
    for child in parent:
        if child.tag == f"{_W}{tag}":
            yield child
        elif child.tag in _WRAPPERS:
            yield from _children(_wrapped_content(child), tag)


def _wrapped_content(wrapper: Any) -> Any:
    # A content control keeps its content in w:sdtContent; custom XML holds it
    # directly.
    content = wrapper.find(f"{_W}sdtContent")
    return wrapper if content is None else content


def _has_ancestor(element: Any, stop: Any, tags: frozenset) -> bool:
    node = element.getparent()
    while node is not None and node is not stop:
        if node.tag in tags:
            return True
        node = node.getparent()
    return False


def load_from_xml_v2(base_uri, rels_item_xml):
    """Return |_SerializedRelationships| instance loaded with the relationships.

    contained in *rels_item_xml*.collection if *rels_item_xml* is |None|.
    """
    srels = _SerializedRelationships()
    if rels_item_xml is not None:
        rels_elm = parse_xml(rels_item_xml)
        for rel_elm in rels_elm.Relationship_lst:
            if rel_elm.target_ref in ("../NULL", "NULL"):
                continue
            srels._srels.append(_SerializedRelationship(base_uri, rel_elm))
    return srels


class DocxKnowledge(Knowledge):
    """Docx Knowledge."""

    def __init__(
        self,
        file_path: Optional[str] = None,
        knowledge_type: Any = KnowledgeType.DOCUMENT,
        encoding: Optional[str] = "utf-8",
        loader: Optional[Any] = None,
        metadata: Optional[Dict[str, Union[str, List[str]]]] = None,
        **kwargs: Any,
    ) -> None:
        """Create Docx Knowledge with Knowledge arguments.

        Args:
            file_path(str,  optional): file path
            knowledge_type(KnowledgeType, optional): knowledge type
            encoding(str, optional): csv encoding
            loader(Any, optional): loader
        """
        super().__init__(
            path=file_path,
            knowledge_type=knowledge_type,
            data_loader=loader,
            metadata=metadata,
            **kwargs,
        )
        self._encoding = encoding

    def _load(self) -> List[Document]:
        """Load docx document from loader."""
        if self._loader:
            documents = self._loader.load()
        else:
            docs = []
            _SerializedRelationships.load_from_xml = load_from_xml_v2  # type: ignore
            doc = docx.Document(self._path)
            content = list(_docx_blocks(doc.element.body))
            metadata = {"source": self._path}
            if self._metadata:
                metadata.update(self._metadata)  # type: ignore
            docs.append(Document(content="\n".join(content), metadata=metadata))
            return docs
        return [Document.langchain2doc(lc_document) for lc_document in documents]

    @classmethod
    def support_chunk_strategy(cls) -> List[ChunkStrategy]:
        """Return support chunk strategy."""
        return [
            ChunkStrategy.CHUNK_BY_SIZE,
            ChunkStrategy.CHUNK_BY_PARAGRAPH,
            ChunkStrategy.CHUNK_BY_SEPARATOR,
        ]

    @classmethod
    def default_chunk_strategy(cls) -> ChunkStrategy:
        """Return default chunk strategy."""
        return ChunkStrategy.CHUNK_BY_SIZE

    @classmethod
    def type(cls) -> KnowledgeType:
        """Return knowledge type."""
        return KnowledgeType.DOCUMENT

    @classmethod
    def document_type(cls) -> DocumentType:
        """Return document type."""
        return DocumentType.DOCX
