import re
import zipfile

import openpyxl
import pytest

from ..excel import ExcelKnowledge


@pytest.fixture
def multi_sheet_xlsx(tmp_path):
    path = tmp_path / "multi_sheet.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["name", "value"])
    ws1.append(["a", 1])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["name", "value"])
    ws2.append(["b", 2])
    wb.save(path)
    return str(path)


@pytest.fixture
def numeric_only_xlsx(tmp_path):
    path = tmp_path / "numeric_only.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NumericSheet"
    for r in range(6):
        ws.append([r, r + 1, r + 2])
    wb.save(path)
    return str(path)


@pytest.fixture
def formula_xlsx(tmp_path):
    """A workbook as Excel saves it: a formula cell also stores its result."""
    raw = tmp_path / "raw.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(["item", "price", "with tax"])
    ws.append(["apple", 3, "=B2*1.2"])
    ws.append(["total", "=SUM(B2:B2)", "=C2"])
    wb.save(raw)
    # openpyxl writes formulas without a result, so add the <v> Excel writes.
    results = {"C2": "3.6", "B3": "3", "C3": "3.6"}
    path = tmp_path / "formulas.xlsx"
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(path, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode()
                for ref, value in results.items():
                    xml, count = re.subn(
                        rf'(<c r="{ref}"[^>]*><f>[^<]*</f>)(<v\s*/>|<v></v>)?',
                        rf"\1<v>{value}</v>",
                        xml,
                    )
                    assert count == 1
                data = xml.encode()
            dst.writestr(item, data)
    return str(path)


def test_load_reads_all_sheets(multi_sheet_xlsx):
    knowledge = ExcelKnowledge(file_path=multi_sheet_xlsx)
    docs = knowledge._load()
    sheets_seen = {doc.metadata["sheet_name"] for doc in docs}
    assert sheets_seen == {"Sheet1", "Sheet2"}


def test_load_does_not_crash_without_header_row(numeric_only_xlsx):
    knowledge = ExcelKnowledge(file_path=numeric_only_xlsx)
    docs = knowledge._load()
    assert len(docs) == 6


def test_load_reads_formula_results_not_formulas(formula_xlsx):
    docs = ExcelKnowledge(file_path=formula_xlsx)._load()
    assert [doc.content for doc in docs] == [
        "item: apple\nprice: 3\nwith tax: 3.6",
        "item: total\nprice: 3\nwith tax: 3.6",
    ]
