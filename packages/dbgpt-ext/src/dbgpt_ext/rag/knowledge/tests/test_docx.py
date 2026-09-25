import docx
from docx.oxml import parse_xml

from ..docx import DocxKnowledge

_NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _save(document, tmp_path) -> str:
    path = str(tmp_path / "test_document.docx")
    document.save(path)
    return path


def _docx_with_body(tmp_path, *fragments: str) -> str:
    """Save a DOCX whose body holds the given WordprocessingML fragments."""
    document = docx.Document()
    section_properties = document.element.body[-1]
    body = parse_xml(f"<w:body {_NAMESPACES}>{''.join(fragments)}</w:body>")
    for element in list(body):
        section_properties.addprevious(element)
    return _save(document, tmp_path)


def _load_text(file_path: str) -> str:
    documents = DocxKnowledge(file_path=file_path)._load()
    assert len(documents) == 1
    return documents[0].content


def test_load_from_docx(tmp_path):
    document = docx.Document()
    document.add_paragraph("This is the first paragraph.")
    document.add_paragraph("This is the second paragraph.")
    file_path = _save(document, tmp_path)

    documents = DocxKnowledge(file_path=file_path)._load()

    assert len(documents) == 1
    assert (
        documents[0].content
        == "This is the first paragraph.\nThis is the second paragraph."
    )
    assert documents[0].metadata["source"] == file_path


def test_plain_paragraphs_load_exactly_as_before(tmp_path):
    document = docx.Document()
    document.add_paragraph("First line\twith a tab")
    document.add_paragraph("")
    document.add_paragraph("Second").add_run().add_break()
    file_path = _save(document, tmp_path)

    expected = "\n".join(p.text for p in docx.Document(file_path).paragraphs)
    assert _load_text(file_path) == expected


def test_tables_load_in_document_order(tmp_path):
    document = docx.Document()
    document.add_paragraph("Before the table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).add_paragraph("Ada")  # after the cell's empty first paragraph
    table.cell(1, 1).text = "Engineer\nLead"  # a line break inside the cell
    document.add_paragraph("After the table")

    assert _load_text(_save(document, tmp_path)) == (
        "Before the table\nName | Role\nAda | Engineer Lead\nAfter the table"
    )


def test_a_text_box_loads_once_after_its_paragraph(tmp_path):
    # Word writes every text box twice: a DrawingML shape and a VML fallback copy.
    box = "<w:txbxContent><w:p>" + _run("Callout text") + "</w:p></w:txbxContent>"
    choice = f"<w:drawing><wps:wsp><wps:txbx>{box}</wps:txbx></wps:wsp></w:drawing>"
    fallback = f"<w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict>"
    file_path = _docx_with_body(
        tmp_path,
        "<w:p>"
        + _run("Host paragraph")
        + "<w:r><mc:AlternateContent>"
        + f'<mc:Choice Requires="wps">{choice}</mc:Choice>'
        + f"<mc:Fallback>{fallback}</mc:Fallback>"
        + "</mc:AlternateContent></w:r></w:p>",
    )

    assert _load_text(file_path) == "Host paragraph\nCallout text"


def test_tracked_insertions_load_and_deletions_do_not(tmp_path):
    file_path = _docx_with_body(
        tmp_path,
        "<w:p>"
        + _run("The fee is ")
        + '<w:del w:id="1" w:author="a">'
        + "<w:r><w:delText>ten</w:delText><w:tab/></w:r></w:del>"
        + f'<w:ins w:id="2" w:author="a">{_run("twelve")}</w:ins>'
        + '<w:moveFrom w:id="3" w:author="a">'
        + "<w:r><w:t> moved away</w:t></w:r></w:moveFrom>"
        + _run(" euros.")
        + "</w:p>",
    )

    assert _load_text(file_path) == "The fee is twelve euros."


def test_content_controls_fields_and_smart_tags_load(tmp_path):
    file_path = _docx_with_body(
        tmp_path,
        "<w:p>"
        + _run("Client: ")
        + f"<w:sdt><w:sdtPr/><w:sdtContent>{_run('Acme Corp')}</w:sdtContent></w:sdt>"
        + "</w:p>",
        "<w:sdt><w:sdtPr/><w:sdtContent>"
        + f"<w:p>{_run('Block control')}</w:p></w:sdtContent></w:sdt>",
        '<w:p><w:fldSimple w:instr=" DOCPROPERTY Company ">'
        + f"{_run('Field result')}</w:fldSimple></w:p>",
        '<w:p><w:smartTag w:uri="urn:x" w:element="place">'
        + f"{_run('Smart tag')}</w:smartTag></w:p>",
        '<w:customXml w:element="clause">'
        + f"<w:p>{_run('Custom XML')}</w:p></w:customXml>",
        "<w:tbl><w:tr><w:tc><w:p>"
        + _run("Row")
        + "</w:p></w:tc></w:tr><w:sdt><w:sdtContent><w:tr><w:tc><w:p>"
        + _run("Repeated row")
        + "</w:p></w:tc></w:tr></w:sdtContent></w:sdt></w:tbl>",
    )

    assert _load_text(file_path) == (
        "Client: Acme Corp\nBlock control\nField result\nSmart tag\nCustom XML\n"
        "Row\nRepeated row"
    )


def test_ruby_base_text_loads_without_its_guide(tmp_path):
    file_path = _docx_with_body(
        tmp_path,
        "<w:p><w:r><w:ruby><w:rubyPr/>"
        f"<w:rt>{_run('kanji')}</w:rt><w:rubyBase>{_run('漢字')}</w:rubyBase>"
        "</w:ruby></w:r></w:p>",
    )

    assert _load_text(file_path) == "漢字"
